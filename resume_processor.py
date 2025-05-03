import os
import pdfplumber
from docx import Document as DocxDocument
from typing import List, Dict

from settings import settings
from langchain_groq import ChatGroq
from langchain_qdrant import QdrantVectorStore
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document
from langchain.docstore.document import Document as LangchainDocument
from qdrant_client import QdrantClient
from qdrant_client.http.models import Distance, VectorParams
import groq

# Initialize Qdrant and Embeddings
try:
    qdrant_client = QdrantClient(url=settings.QDRANT_URL, api_key=settings.QDRANT_API_KEY)
except Exception as e:
    raise ValueError(f"Failed to initialize Qdrant client: {e}")

embedding_model = settings.SENTENCE_TRANSFORMER_MODEL or "sentence-transformers/all-MiniLM-L6-v2"
embeddings = HuggingFaceEmbeddings(model_name=embedding_model)

def sanitize_input(text: str) -> str:
    return text.replace("```", "").replace("{{", "").strip()

def load_resume(file_path: str) -> str:
    if not os.path.exists(file_path):
        raise ValueError(f"File not found: {file_path}")

    try:
        if file_path.endswith(".pdf"):
            with pdfplumber.open(file_path) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        elif file_path.endswith(".txt"):
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        elif file_path.endswith(".docx"):
            doc = DocxDocument(file_path)
            text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
        else:
            raise ValueError("Unsupported file type. Use PDF, TXT, or DOCX.")
        return sanitize_input(text.strip())
    except Exception as e:
        raise ValueError(f"Error loading resume: {e}")

def semantic_chunk_text(text: str, chunk_size=500, chunk_overlap=50) -> List[str]:
    try:
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ".", " "]
        )
        return splitter.split_text(sanitize_input(text))
    except Exception as e:
        raise ValueError(f"Error chunking text: {e}")

def store_resume_chunks_in_qdrant(chunks: List[str], metadata: dict = None) -> QdrantVectorStore:
    try:
        if metadata is None:
            metadata = {}

        # Wrap chunks in LangChain Documents
        docs = [LangchainDocument(page_content=chunk, metadata=metadata) for chunk in chunks]

        # Create or connect to a Qdrant collection using LangChain
        vector_store = QdrantVectorStore.from_documents(
            documents=docs,
            embedding=embeddings,
            url=settings.QDRANT_URL,
            api_key=settings.QDRANT_API_KEY,
            collection_name=settings.QDRANT_COLLECTION_NAME
        )
        return vector_store
    except Exception as e:
        print("❌ Failed to store resume in Qdrant using LangChain:", e)
        raise

def validate_resume_with_llm(resume_text: str, job_prompt: str = None) -> str:
    prompt = f"""
You are an expert resume reviewer. Evaluate the following resume for structure, completeness, and relevance.
Highlight any missing sections (e.g., Contact Info, Summary, Skills, Experience, Education), and provide clear suggestions for improvement.
{"If a job prompt is provided, also evaluate how well the resume aligns with the job requirements: " + sanitize_input(job_prompt) if job_prompt else ""}

Resume Content:
{sanitize_input(resume_text)}

Output format:
- ✅ [Good sections]
- ⚠️ [Areas needing improvement]
- ❗ [Critical missing items]
- Suggestions:
"""

    try:
        groq_client = groq.Client(api_key=settings.GROQ_KEY)
        response = groq_client.chat.completions.create(
            model=settings.GROQ_MODEL,
            messages=[
                {"role": "system", "content": "You are a resume reviewer AI."},
                {"role": "user", "content": prompt}
            ]
        )
        return response.choices[0].message.content.strip()
    except groq.APIError as e:
        raise ValueError(f"Groq API error: {e}")
    except Exception as e:
        raise ValueError(f"Error validating resume: {e}")

def evaluate_resume_for_roles(resume_text: str, job_roles: List[Dict], vector_store: QdrantVectorStore) -> List[Dict]:
    results = []

    try:
        groq_client = groq.Client(api_key=settings.GROQ_KEY)
    except Exception as e:
        raise ValueError(f"Failed to initialize Groq client: {e}")

    for role in job_roles:
        role_name = sanitize_input(role.get("role_name", "Unnamed Role"))
        job_prompt = sanitize_input(role.get("description", ""))
        required_skills = [sanitize_input(skill) for skill in role.get("skills", [])]
        required_experience = role.get("experience", 0)
        required_education = sanitize_input(role.get("education", ""))

        try:
            query_embedding = embeddings.embed_query(job_prompt)
            search_results = vector_store.search(query_embedding, k=5)
            relevant_chunks = "\n".join([res.page_content for res in search_results])
        except Exception as e:
            relevant_chunks = f"Error in vector search: {e}"

        prompt = f"""
You are an expert hiring manager. Evaluate the following resume to determine its eligibility for the job role described below. Assign a match score (0-100) based on:
- Presence of required skills (50%)
- Relevant experience (30%)
- Education and certifications (20%)

Job Role: {role_name}
Description: {job_prompt}
Required Skills: {', '.join(required_skills)}
Required Experience: {required_experience} years
Required Education: {required_education}

Relevant Resume Sections:
{relevant_chunks}

Resume Content:
{sanitize_input(resume_text)}

Output format:
- Match Score: [0-100]
- Skills Match: [Details]
- Experience Match: [Details]
- Education Match: [Details]
- Feedback: [Eligibility and suggestions]
"""

        try:
            response = groq_client.chat.completions.create(
                model=settings.GROQ_MODEL,
                messages=[
                    {"role": "system", "content": "You are a hiring manager AI."},
                    {"role": "user", "content": prompt}
                ]
            )
            result = response.choices[0].message.content.strip()
        except groq.APIError as e:
            result = f"Error evaluating role {role_name}: {e}"
        except Exception as e:
            result = f"Unexpected error evaluating role {role_name}: {e}"

        results.append({
            "role_name": role_name,
            "evaluation": result
        })

    return results